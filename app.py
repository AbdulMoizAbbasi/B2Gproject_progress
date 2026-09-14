from flask import Flask, jsonify
from flask_cors import CORS
from openpyxl import load_workbook

from update_news import (
    EXCEL_FILE,
    PROJECT_SHEET,
    NEWS_SHEET,
    refresh_project
)


app = Flask(__name__)


# CORS configuration
CORS(
    app,
    resources={
        r"/api/*": {
            "origins": [
                "http://localhost:5500",
                "http://127.0.0.1:5500",
                "http://localhost:5173",
                "http://127.0.0.1:5173",
                "http://localhost:5174",
                "http://127.0.0.1:5174"
            ]
        }
    }
)


def get_projects():

    wb = load_workbook(
        EXCEL_FILE,
        data_only=True
    )

    ws = wb[PROJECT_SHEET]

    headers = {}

    for cell in ws[1]:

        if cell.value:
            headers[str(cell.value).strip()] = cell.column

    project_id_col = headers["Project ID"]
    project_name_col = headers["Project / Scheme Name"]

    projects = []

    for row in range(2, ws.max_row + 1):

        project_id = ws.cell(
            row,
            project_id_col
        ).value

        project_name = ws.cell(
            row,
            project_name_col
        ).value

        if not project_name:
            continue

        projects.append({
            "project_id": project_id,
            "project_name": str(project_name).strip()
        })

    return projects


def get_project_details(project_id):

    wb = load_workbook(
        EXCEL_FILE,
        data_only=True
    )

    ws = wb[PROJECT_SHEET]

    headers = {}

    for cell in ws[1]:

        if cell.value:
            headers[str(cell.value).strip()] = cell.column

    project_id_col = headers["Project ID"]

    for row in range(2, ws.max_row + 1):

        current_id = ws.cell(
            row,
            project_id_col
        ).value

        if str(current_id).strip() == str(project_id).strip():

            project = {}

            for name, column in headers.items():

                value = ws.cell(
                    row,
                    column
                ).value

                if hasattr(value, "isoformat"):
                    value = value.isoformat()

                project[name] = value

            return project

    return None


def get_project_news(project_id):

    wb = load_workbook(
        EXCEL_FILE,
        data_only=True
    )

    if NEWS_SHEET not in wb.sheetnames:
        return []

    ws = wb[NEWS_SHEET]

    headers = {}

    for cell in ws[1]:

        if cell.value:
            headers[str(cell.value).strip()] = cell.column

    if "Project ID" not in headers:
        return []

    news = []

    for row in range(2, ws.max_row + 1):

        current_id = ws.cell(
            row,
            headers["Project ID"]
        ).value

        if str(current_id).strip() != str(project_id).strip():
            continue

        item = {}

        for name, column in headers.items():

            value = ws.cell(
                row,
                column
            ).value

            if hasattr(value, "isoformat"):
                value = value.isoformat()

            item[name] = value

        news.append(item)

    return news


@app.route("/")
def home():

    return jsonify({
        "message": "Project Tracking API is running"
    })


@app.route("/api/projects")
def projects():

    try:

        return jsonify(
            get_projects()
        )

    except Exception as e:

        return jsonify({
            "error": str(e)
        }), 500


@app.route("/api/projects/<project_id>")
def project(project_id):

    try:

        project_data = get_project_details(
            project_id
        )

        if not project_data:

            return jsonify({
                "error": "Project not found"
            }), 404

        return jsonify(project_data)

    except Exception as e:

        return jsonify({
            "error": str(e)
        }), 500


@app.route("/api/projects/<project_id>/news")
def project_news(project_id):

    try:

        project_data = get_project_details(
            project_id
        )

        if not project_data:

            return jsonify({
                "error": "Project not found"
            }), 404

        news = get_project_news(
            project_id
        )

        return jsonify({
            "project_id": project_id,
            "project_name": project_data.get(
                "Project / Scheme Name"
            ),
            "news": news
        })

    except Exception as e:

        return jsonify({
            "success": False,
            "error": str(e)
        }), 500


@app.route(
    "/api/projects/<project_id>/news/refresh",
    methods=["POST"]
)
def refresh_news(project_id):

    try:

        result = refresh_project(
            project_id=project_id
        )

        news = get_project_news(
            project_id
        )

        return jsonify({
            "success": True,
            "project_id": project_id,
            "project_name": result["project_name"],
            "new_news": result["new_news"],
            "news": news
        })

    except Exception as e:

        print(
            f"Refresh error for project {project_id}: {e}"
        )

        return jsonify({
            "success": False,
            "error": str(e)
        }), 500


if __name__ == "__main__":

    app.run(
        debug=True,
        host="0.0.0.0",
        port=5000
    )