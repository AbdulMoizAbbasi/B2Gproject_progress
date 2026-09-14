import os
from datetime import datetime
from urllib.parse import urlparse

from dotenv import load_dotenv
from openpyxl import load_workbook
from tavily import TavilyClient
from groq import Groq


EXCEL_FILE = "Consolidated_ICT_Projects_2026-27_1.xlsx"
PROJECT_SHEET = "ICT Projects - Consolidated"
NEWS_SHEET = "Project News"

MAX_RESULTS = 3
SEARCH_DEPTH = "advanced"


load_dotenv()

TAVILY_API_KEY = os.getenv("api_key")
GROQ_API_KEY = os.getenv("groq_api_key")

if not TAVILY_API_KEY:
    raise ValueError("Tavily API key not found in .env")

if not GROQ_API_KEY:
    raise ValueError("Groq API key not found in .env")

tavily = TavilyClient(api_key=TAVILY_API_KEY)
groq = Groq(api_key=GROQ_API_KEY)


def get_source(url):
    try:
        return urlparse(url).netloc.replace("www.", "")
    except:
        return ""


def parse_date(date_value):
    if not date_value:
        return datetime.min

    if isinstance(date_value, datetime):
        return date_value

    try:
        return datetime.fromisoformat(
            str(date_value).replace("Z", "+00:00")
        ).replace(tzinfo=None)
    except:
        return datetime.min


def get_project_column(ws, column_name):
    for cell in ws[1]:
        if str(cell.value).strip() == column_name:
            return cell.column

    raise ValueError(f"Column '{column_name}' not found")


def create_news_sheet(wb):
    if NEWS_SHEET in wb.sheetnames:
        ws = wb[NEWS_SHEET]
    else:
        ws = wb.create_sheet(NEWS_SHEET)

        headers = [
            "Project ID",
            "Project Name",
            "News Date",
            "News Title",
            "Keypoints",
            "Source",
            "URL",
            "Last Updated"
        ]

        ws.append(headers)

    return ws


def get_existing_urls(news_ws, project_id):
    urls = set()

    for row in news_ws.iter_rows(min_row=2, values_only=True):
        if str(row[0]).strip() == str(project_id).strip() and row[6]:
            urls.add(str(row[6]).strip())

    return urls


def generate_keypoints(title, content):
    prompt = f"""
        Summarize this news article into EXACTLY 2 short bullet points.

        Title:
        {title}

        Article:
        {content}

        Rules:
        - Each bullet must be maximum 15 words.
        - Focus only on NEW and IMPORTANT information.
        - Focus on project progress, funding, tenders, implementation,
        decisions, timelines, or milestones.
        - Do not copy long sentences from the article.
        - Do not include background information unless essential.
        - Do not include emojis, hashtags, URLs, or source names.
        - Do not repeat the project name unnecessarily.
        - If the article contains no meaningful project update, summarize
        the most relevant information briefly.
        - Return ONLY 2 bullet points.
        """

    response = groq.chat.completions.create(
        model="openai/gpt-oss-20b",
        messages=[
            {
                "role": "user",
                "content": prompt
            }
        ],
        temperature=0.1
    )

    return response.choices[0].message.content.strip()


def refresh_project(project_id=None, project_name=None):

    wb = load_workbook(EXCEL_FILE)

    if PROJECT_SHEET not in wb.sheetnames:
        raise ValueError(
            f"Sheet '{PROJECT_SHEET}' not found"
        )

    projects_ws = wb[PROJECT_SHEET]
    news_ws = create_news_sheet(wb)

    project_id_col = get_project_column(
        projects_ws,
        "Project ID"
    )

    project_name_col = get_project_column(
        projects_ws,
        "Project / Scheme Name"
    )

    selected_project_id = None
    selected_project_name = None

    for row in range(2, projects_ws.max_row + 1):

        current_id = projects_ws.cell(
            row,
            project_id_col
        ).value

        current_name = projects_ws.cell(
            row,
            project_name_col
        ).value

        if project_id is not None:

            if str(current_id).strip() == str(project_id).strip():
                selected_project_id = current_id
                selected_project_name = str(current_name).strip()
                break

        elif project_name is not None:

            if str(current_name).strip() == str(project_name).strip():
                selected_project_id = current_id
                selected_project_name = str(current_name).strip()
                break

    if selected_project_name is None:
        raise ValueError("Project not found")

    print(f"\nSearching:")
    print(selected_project_name)

    response = tavily.search(
        query=selected_project_name,
        search_depth=SEARCH_DEPTH,
        max_results=MAX_RESULTS,
        include_answer=False,
        include_raw_content=False
    )

    existing_urls = get_existing_urls(
        news_ws,
        selected_project_id
    )

    new_count = 0

    for result in response.get("results", []):

        url = result.get("url", "").strip()

        if not url:
            continue

        if url in existing_urls:
            continue

        title = result.get("title", "").strip()
        content = result.get("content", "").strip()
        published_date = result.get("published_date")
        source = get_source(url)

        try:
            keypoints = generate_keypoints(
                title,
                content
            )
        except Exception as e:
            print(f"Groq error: {e}")
            keypoints = content

        news_ws.append([
            selected_project_id,
            selected_project_name,
            published_date,
            title,
            keypoints,
            source,
            url,
            datetime.now()
        ])

        existing_urls.add(url)
        new_count += 1

    rows = list(
        news_ws.iter_rows(
            min_row=2,
            values_only=True
        )
    )

    rows.sort(
        key=lambda row: parse_date(row[2]),
        reverse=True
    )

    if news_ws.max_row > 1:
        news_ws.delete_rows(
            2,
            news_ws.max_row - 1
        )

    for row in rows:
        news_ws.append(row)

    news_ws.freeze_panes = "A2"
    news_ws.auto_filter.ref = news_ws.dimensions

    widths = {
        "A": 15,
        "B": 60,
        "C": 20,
        "D": 70,
        "E": 100,
        "F": 30,
        "G": 100,
        "H": 20
    }

    for column, width in widths.items():
        news_ws.column_dimensions[column].width = width

    wb.save(EXCEL_FILE)

    print(f"New news added: {new_count}")

    return {
        "project_id": selected_project_id,
        "project_name": selected_project_name,
        "new_news": new_count
    }


def get_all_projects():

    wb = load_workbook(EXCEL_FILE)

    ws = wb[PROJECT_SHEET]

    project_id_col = get_project_column(
        ws,
        "Project ID"
    )

    project_name_col = get_project_column(
        ws,
        "Project / Scheme Name"
    )

    projects = []

    for row in range(2, ws.max_row + 1):

        project_id = ws.cell(
            row,
            project_id_col
        ).value

        project_name = ws.cell(
            row,
            project_name_col
        ).value

        if project_name:

            projects.append({
                "project_id": project_id,
                "project_name": str(project_name).strip()
            })

    return projects


if __name__ == "__main__":

    projects = get_all_projects()

    for project in projects:

        try:

            refresh_project(
                project_id=project["project_id"]
            )

        except Exception as e:

            print(
                f"Error for {project['project_name']}: {e}"
            )
